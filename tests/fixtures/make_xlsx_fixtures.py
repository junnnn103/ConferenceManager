"""테스트용 합성 엑셀을 만든다. 원본 사내 엑셀은 커밋되지 않으므로
같은 시트 구조를 가진 축약본이 필요하다.

실행: python tests/fixtures/make_xlsx_fixtures.py
"""

from pathlib import Path

import openpyxl

HERE = Path(__file__).parent


def make_grade_xlsx(path: Path) -> None:
    """26년 우수 학회 List.xlsx 와 같은 구조.
    B열=No, C열=등급, D열=약어, E열=Full Name, 데이터는 5행부터.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["B2"] = "26년 최우수/우수 학회 List"
    ws["B4"], ws["C4"], ws["D4"], ws["E4"] = "No", "등급", "약어", "Full Name"
    rows = [
        (1, "최우수", "cvpr", "Computer Vision and Pattern Recognition"),
        (2, "최우수", "nips", "Neural Information Processing Systems"),
        (3, "최우수", "iccv/eccv", "International Conference on Computer Vision / European Conference on Computer Vision"),
        (4, "우수", "wacv", "IEEE Winter Conference on Applications of Computer Vision"),
        (5, "우수", "SID", "SID DISPLAYWEEK"),
    ]
    for i, row in enumerate(rows, start=5):
        for j, value in enumerate(row):
            ws.cell(row=i, column=2 + j, value=value)
    wb.save(path)


def make_ai_specialist_xlsx(path: Path) -> None:
    """260406_AI_Specialist_인정학회리스트.xlsx 와 같은 구조.
    B열=대분류(첫 행에만), C열=No, D열=약어, E열=학회명, 데이터는 5행부터.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["B2"] = "AI Specialist 인정 학회 List"
    ws["B4"], ws["C4"], ws["D4"], ws["E4"] = "대분류", "No", "약어", "학회명(영문)"
    rows = [
        ("AI/Data", 1, "CVPR", "Computer Vision and Pattern Recognition"),
        (None, 2, "NeurIPS", "Neural Information Processing Systems"),
        (None, 3, "ICCV", "International Conference on Computer Vision"),
        (None, 4, "ECCV", "European Conference on Computer Vision"),
        (None, 5, "WACV", "IEEE Winter Conference on Applications of Computer Vision"),
    ]
    for i, row in enumerate(rows, start=5):
        for j, value in enumerate(row):
            if value is not None:
                ws.cell(row=i, column=2 + j, value=value)
    wb.save(path)


if __name__ == "__main__":
    make_grade_xlsx(HERE / "grades.xlsx")
    make_ai_specialist_xlsx(HERE / "ai_specialist.xlsx")
    print("wrote", HERE / "grades.xlsx", "and", HERE / "ai_specialist.xlsx")
